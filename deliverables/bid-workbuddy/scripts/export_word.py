#!/usr/bin/env python3
"""从 build/book.json 导出可更新目录的中文 DOCX；不支持的 Markdown 拒绝导出。"""

import argparse
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
from urllib.parse import unquote, urlsplit

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


DEFAULT_PROFILE = {
    'body_font': '宋体', 'heading_font': '黑体', 'body_size_pt': 12,
    'line_spacing': 1.5, 'margin_top_cm': 2.5, 'margin_bottom_cm': 2.5,
    'margin_left_cm': 2.8, 'margin_right_cm': 2.5, 'toc_levels': 3,
}
HEADING = re.compile(r'^(#{1,6})\s+(.+?)\s*#*\s*$')
LIST = re.compile(r'^( *)([-+*]|\d+[.)])\s+(.+)$')
IMAGE = re.compile(r'^!\[([^\]]+)\]\(([^)]+)\)$')
INLINE = re.compile(r'\*\*([^*\n]+)\*\*|\*([^*\n]+)\*|\[([^\]\n]+)\]\(([^)\s]+)\)|\\([\\`*_{}\[\]()#+.!|>~-])')


def require(condition, message):
    """在输入错误时提供可直接修正的诊断。"""
    if not condition:
        raise ValueError(message)


def local_file(workspace, value):
    """解析工作区内文件，拒绝 URL、绝对路径和经链接逃逸的路径。"""
    require(isinstance(value, str) and value.strip(), '图片路径不能为空')
    decoded = unquote(value)
    require(not urlsplit(decoded).scheme and not decoded.startswith(('/', '\\'))
            and not re.match(r'^[A-Za-z]:', decoded), f'图片必须是工作区相对路径：{value}')
    path = (workspace / decoded.replace('\\', '/')).resolve()
    require(path.is_relative_to(workspace), f'图片路径越过工作区：{value}')
    require(path.is_file(), f'图片不存在：{value}')
    require(path.suffix.lower() in ('.png', '.jpg', '.jpeg', '.bmp', '.gif', '.tif', '.tiff'),
            f'图片格式不支持，请先转换为 PNG：{value}')
    return path


def load_profile(path):
    """读取真实排版参数；拒绝未知参数及无效尺寸。"""
    values = {} if path is None else json.loads(Path(path).read_text(encoding='utf-8-sig'))
    require(isinstance(values, dict), '排版配置必须为 JSON 对象')
    require(not set(values) - DEFAULT_PROFILE.keys(), f'未知排版参数：{sorted(set(values) - DEFAULT_PROFILE.keys())}')
    profile = {**DEFAULT_PROFILE, **values}
    for name in ('body_font', 'heading_font'):
        require(isinstance(profile[name], str) and profile[name].strip(), f'{name} 必须为字体名称')
    for name, lower, upper in [('body_size_pt', 9, 24), ('line_spacing', 1, 3)] + [
            (f'margin_{side}_cm', 1, 5) for side in ('top', 'bottom', 'left', 'right')]:
        number = profile[name]
        require(type(number) in (int, float) and math.isfinite(number) and lower <= number <= upper,
                f'{name} 必须在 {lower} 至 {upper} 之间')
    require(type(profile['toc_levels']) is int and 1 <= profile['toc_levels'] <= 6, 'toc_levels 必须为 1 至 6')
    return profile


def load_book(workspace, draft):
    """读取已编排正文；核验目录层级及程序编号，正式导出要求 ready。"""
    book = json.loads((workspace / 'build' / 'book.json').read_text(encoding='utf-8-sig'))
    require(isinstance(book, dict), 'book.json 必须为对象')
    require(book.get('status') in ('ready', 'draft'), 'book.status 必须为 ready 或 draft')
    require(draft or book['status'] == 'ready', '正文尚未 ready；修复审核问题，或明确使用 --draft')
    require(isinstance(book.get('issues'), list), 'book.issues 必须为数组')
    require(isinstance(book.get('title'), str) and book['title'].strip(), 'book.title 必须为非空文字')
    for name in ('bidder', 'date'):
        require(isinstance(book.get(name), str), f'book.{name} 必须为文字；未知时使用空字符串')
    require(isinstance(book.get('sections'), list) and book['sections'], 'book.sections 不能为空')
    counters = [0] * 6
    previous = 0
    for index, section in enumerate(book['sections']):
        require(isinstance(section, dict), f'章节 {index + 1} 必须为对象')
        level = section.get('level')
        require(type(level) is int and 1 <= level <= 6 and level <= previous + 1,
                f'章节 {index + 1} 层级必须从 1 开始且不能跳级，最大为 6')
        for key in ('title', 'number', 'body'):
            require(isinstance(section.get(key), str), f'章节 {index + 1} 的 {key} 必须为字符串')
        require(section['title'].strip(), f'章节 {index + 1} 标题不能为空')
        counters[level - 1] += 1
        counters[level:] = [0] * (6 - level)
        expected = '.'.join(map(str, counters[:level]))
        require(section['number'] == expected, f'章节 {index + 1} 编号应为 {expected}，请重新编排 book.json')
        previous = level
    return book


def element(name, **attributes):
    """构造带 Word 命名空间的元素。"""
    node = OxmlElement(f'w:{name}')
    for key, value in attributes.items():
        node.set(qn(f'w:{key}'), str(value))
    return node


def field(paragraph, instruction, placeholder=''):
    """写入真实 Word 域，不估算或编造页码。"""
    paragraph._p.append(element('r'))
    paragraph._p[-1].append(element('fldChar', fldCharType='begin', dirty='true'))
    run = element('r')
    text = element('instrText')
    text.set(qn('xml:space'), 'preserve')
    text.text = f' {instruction} '
    run.append(text)
    paragraph._p.append(run)
    run = element('r')
    run.append(element('fldChar', fldCharType='separate'))
    paragraph._p.append(run)
    if placeholder:
        paragraph.add_run(placeholder)
    run = element('r')
    run.append(element('fldChar', fldCharType='end'))
    paragraph._p.append(run)


def toc_instructions(document):
    """读取简单域和跨 Run 的复合目录域，不把域结果当作指令。"""
    instructions, stack = [], []
    for node in document.element.iter():
        if node.tag == qn('w:fldSimple'):
            instructions.append(node.get(qn('w:instr'), ''))
        elif node.tag == qn('w:fldChar'):
            kind = node.get(qn('w:fldCharType'))
            if kind == 'begin':
                stack.append([])
            elif kind in ('separate', 'end') and stack:
                if stack[-1] is not None:
                    instructions.append(''.join(stack[-1]))
                    stack[-1] = None
                if kind == 'end':
                    stack.pop()
        elif node.tag == qn('w:instrText') and stack and stack[-1] is not None:
            stack[-1].append(node.text or '')
    return [code for code in instructions if re.match(r'^\s*TOC(?:\s|$)', code, re.IGNORECASE)]


def validate_template(document, directories):
    """拒绝会改变正文可用宽度或遗漏新增大纲标题的模板。"""
    for section in document.sections:
        for columns in section._sectPr.findall(qn('w:cols')):
            count = columns.get(qn('w:num'), '1')
            require(count.isdigit() and int(count) == 1 and len(columns.findall(qn('w:col'))) <= 1,
                    '多栏模板暂不支持；请使用单栏正文模板，避免表格和图片越过栏宽')
    for code in directories:
        switches = re.sub(r'"[^"]*"', '', code)
        require(re.search(r'\\u(?:\s|$)', switches, re.IGNORECASE),
                '模板原有目录未启用大纲级别，无法确认收录新增标题；请在 Word 目录选项中启用大纲级别（TOC \\u），保留其他原有设置后重试')


def font(style, name, size, bold=False):
    """固定中西文字体、字号和黑色，避免主题色覆盖。"""
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), name)


def configure(document, profile, template):
    """使用专用样式，模板已有样式与页面设置保持原样。"""
    owned_styles = ['Bid Body', 'Bid Cover', 'Bid Caption', 'Bid TOC Title'] + [f'Bid Heading {n}' for n in range(1, 7)]
    if template:
        require(not any(name in document.styles for name in owned_styles),
                '模板含导出器专用 Bid 样式；请使用原始模板并设置正文锚点，避免改变模板已有正文')
    for name, size, bold in [('Bid Body', profile['body_size_pt'], False), ('Bid Cover', 26, True),
                              ('Bid Caption', 10.5, False), ('Bid TOC Title', 18, True)]:
        style = document.styles[name] if name in document.styles else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles['Normal']
        font(style, profile['heading_font'] if bold else profile['body_font'], size, bold)
        style.paragraph_format.line_spacing = profile['line_spacing']
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.widow_control = True
    document.styles['Bid Cover'].base_style = document.styles['Title']
    for level in range(1, 7):
        name = f'Bid Heading {level}'
        style = document.styles[name] if name in document.styles else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles[f'Heading {level}']
        font(style, profile['heading_font'], [18, 16, 14, 13, 12, 12][level - 1], True)
        style.paragraph_format.space_before = Pt(14 if level < 3 else 9)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        properties = style.element.get_or_add_pPr()
        properties.append(element('outlineLvl', val=level - 1))
        properties.append(element('numPr'))
        properties[-1].append(element('numId', val=0))
    if not template:
        for section in document.sections:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            for side in ('top', 'bottom', 'left', 'right'):
                setattr(section, f'{side}_margin', Cm(profile[f'margin_{side}_cm']))
            section.different_first_page_header_footer = True
            paragraph = section.footer.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run('第 ')
            field(paragraph, 'PAGE')
            paragraph.add_run(' 页')
    settings = document.settings.element
    old = settings.find(qn('w:updateFields'))
    if old is not None:
        settings.remove(old)
    settings.append(element('updateFields', val='true'))


def inline(paragraph, text):
    """写入粗体、斜体、超链接和转义；拒绝无法保真的内联语法。"""
    cursor = 0
    for match in INLINE.finditer(text):
        plain = text[cursor:match.start()]
        reject_inline(plain)
        paragraph.add_run(plain)
        bold, italic, label, url, escaped = match.groups()
        if bold is not None or italic is not None:
            value = bold if bold is not None else italic
            reject_inline(value)
            run = paragraph.add_run(value)
            run.bold, run.italic = bold is not None, italic is not None
        elif label is not None:
            require(urlsplit(url).scheme in ('http', 'https', 'mailto'), f'链接协议不支持：{url}')
            reject_inline(label)
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
            run = element('r')
            node = element('t')
            node.text = label
            run.append(node)
            link.append(run)
            paragraph._p.append(link)
        else:
            paragraph.add_run(escaped)
        cursor = match.end()
    plain = text[cursor:]
    reject_inline(plain)
    paragraph.add_run(plain)


def reject_inline(text):
    """检测未被支持语法消费的 Markdown 标记。"""
    require(not re.search(r'[`*]|~~|!?\[[^\]]*\](?:\(|\[)|\[\^|<[^>]+>|\$\S[^$]*\$|(?:^|\s)_{1,2}\S', text),
            f'不支持或未闭合的 Markdown 内联语法：{text[:100]}')


def split_row(line):
    """拆分 Markdown 表格单元格并保留转义竖线。"""
    require(line.startswith('|') and line.endswith('|'), '表格每行必须以 | 开始和结束')
    return [part.strip() for part in re.split(r'(?<!\\)\|', line[1:-1])]


def list_numbering(document, ordered, start):
    """建立独立列表计数，保留各列表的起始数字。"""
    root = document.part.numbering_part.element
    abstract_id = max([int(n.get(qn('w:abstractNumId'))) for n in root.findall(qn('w:abstractNum'))] + [-1]) + 1
    num_id = max([int(n.get(qn('w:numId'))) for n in root.findall(qn('w:num'))] + [0]) + 1
    abstract = element('abstractNum', abstractNumId=abstract_id)
    abstract.append(element('multiLevelType', val='multilevel'))
    for depth in range(6):
        level = element('lvl', ilvl=depth)
        level.append(element('start', val=start if depth == 0 else 1))
        level.append(element('numFmt', val='decimal' if ordered else 'bullet'))
        level.append(element('lvlText', val=f'%{depth + 1}.' if ordered else '•'))
        properties = element('pPr')
        properties.append(element('ind', left=360 * (depth + 1), hanging=240))
        level.append(properties)
        abstract.append(level)
    root.append(abstract)
    number = element('num', numId=num_id)
    number.append(element('abstractNumId', val=abstract_id))
    root.append(number)
    return num_id


class Writer:
    """将正文块插入模板锚点之前，或追加到默认文档。"""

    def __init__(self, document, workspace, profile, anchor=None):
        self.document, self.workspace, self.profile, self.anchor = document, workspace, profile, anchor
        section = document.sections[0]
        if anchor is not None:
            for node in anchor._p.itersiblings():
                section_properties = node.find(qn('w:pPr'))
                if node.tag == qn('w:sectPr'):
                    section_properties = node
                elif section_properties is not None:
                    section_properties = section_properties.find(qn('w:sectPr'))
                if section_properties is not None and section_properties.tag == qn('w:sectPr'):
                    from docx.section import Section
                    section = Section(section_properties, document.part)
                    break
        self.width = section.page_width - section.left_margin - section.right_margin
        self.height = section.page_height - section.top_margin - section.bottom_margin
        require(self.width > Cm(3) and self.height > Cm(3), '模板正文可用页面尺寸不足')
        self.tables, self.images = 0, 0

    def place(self, block):
        if self.anchor is not None:
            self.anchor._p.addprevious(block._element)
        return block

    def paragraph(self, text='', style='Bid Body'):
        paragraph = self.place(self.document.add_paragraph(style=style))
        inline(paragraph, text)
        return paragraph

    def heading(self, number, title, level):
        paragraph = self.paragraph(f'{number} {title}', f'Bid Heading {level}')
        return paragraph

    def table(self, rows, caption, alignments):
        self.tables += 1
        paragraph = self.paragraph(f'表 {self.tables}' + (f' {caption}' if caption else ''), 'Bid Caption')
        paragraph.paragraph_format.keep_with_next = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        table = self.place(self.document.add_table(rows=1, cols=len(rows[0])))
        table.autofit = False
        widths = [max(4, min(36, max(len(row[col]) for row in rows))) for col in range(len(rows[0]))]
        total = sum(widths)
        for column, weight in zip(table.columns, widths):
            column.width = int(self.width * weight / total)
        properties = table._tbl.tblPr
        borders = element('tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            borders.append(element(edge, val='single', sz=4, color='D9D9D9'))
        properties.append(borders)
        margins = element('tblCellMar')
        for side in ('top', 'left', 'bottom', 'right'):
            margins.append(element(side, w=90, type='dxa'))
        properties.append(margins)
        for index, values in enumerate(rows):
            row = table.rows[0] if index == 0 else table.add_row()
            if index == 0:
                row._tr.get_or_add_trPr().append(element('tblHeader', val='true'))
            for col, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = int(self.width * widths[col] / total)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                paragraph.style = self.document.styles['Bid Body']
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                inline(paragraph, value)
                if index == 0:
                    cell._tc.get_or_add_tcPr().append(element('shd', fill='E8EDF2'))
                    for run in paragraph.runs:
                        run.bold = True
                if alignments[col] is not None:
                    paragraph.alignment = alignments[col]
                elif len(value) <= 8:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def image(self, label, relative):
        path = local_file(self.workspace, relative)
        from docx.image.image import Image
        info = Image.from_file(str(path))
        scale = min(self.width / info.width, (self.height - Cm(2)) / info.height, 1)
        paragraph = self.paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        picture = paragraph.add_run().add_picture(str(path), width=int(info.width * scale))
        picture._inline.docPr.set('descr', label)
        self.images += 1
        caption = self.paragraph(f'图 {self.images} {label}', 'Bid Caption')
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def markdown(self, section, has_children):
        lines, index, local_counters, previous_heading = section['body'].splitlines(), 0, [0] * 6, 0
        list_id, list_kind, list_depth, expected_numbers = None, None, 0, {}
        while index < len(lines):
            raw, line = lines[index], lines[index].strip()
            index += 1
            if not line:
                list_id = None
                continue
            require('\t' not in raw, 'Markdown 不支持制表符缩进，请使用两个空格一级的列表')
            require(not raw.endswith(('  ', '\\')), 'Markdown 硬换行暂不支持，请改为独立段落')
            image = IMAGE.fullmatch(line)
            heading = HEADING.fullmatch(line)
            item = LIST.fullmatch(raw)
            if heading:
                require(not has_children, f'{section["number"]} 有正式下级章节，正文不能再建立标题')
                relative_level = len(heading[1])
                level = section['level'] + relative_level
                require(level <= 6 and relative_level <= previous_heading + 1,
                        f'{section["number"]} 正文标题跳级或总层级超过 6')
                local_counters[relative_level - 1] += 1
                local_counters[relative_level:] = [0] * (6 - relative_level)
                number = section['number'] + '.' + '.'.join(map(str, local_counters[:relative_level]))
                self.heading(number, heading[2], level)
                previous_heading = relative_level
            elif image:
                self.image(image[1], image[2])
            elif line.startswith('|') or line.startswith('表：'):
                caption = line[2:].strip() if line.startswith('表：') else ''
                if caption:
                    require(index < len(lines), '表题后缺少 Markdown 表格')
                    line = lines[index].strip()
                    index += 1
                rows = [split_row(line)]
                require(index < len(lines), 'Markdown 表格缺少分隔行')
                separator = split_row(lines[index].strip())
                index += 1
                require(len(separator) == len(rows[0]) and all(re.fullmatch(r':?-{3,}:?', v) for v in separator),
                        'Markdown 表格第二行必须为每列至少三个连字符')
                alignments = [(WD_ALIGN_PARAGRAPH.CENTER if value.startswith(':') and value.endswith(':')
                               else WD_ALIGN_PARAGRAPH.RIGHT if value.endswith(':')
                               else WD_ALIGN_PARAGRAPH.LEFT if value.startswith(':') else None) for value in separator]
                while index < len(lines) and lines[index].strip().startswith('|'):
                    row = split_row(lines[index].strip())
                    require(len(row) == len(rows[0]), 'Markdown 表格各行列数必须相同')
                    rows.append(row)
                    index += 1
                self.table(rows, caption, alignments)
            elif item:
                indent, marker, text = item.groups()
                depth, ordered = len(indent) // 2, marker[0].isdigit()
                require(len(indent) % 2 == 0 and depth < 6, '列表缩进必须为两个空格一级，最多 6 级')
                require(not re.match(r'\[[ xX]\]\s', text), '任务复选框不支持，请使用普通列表或表格')
                if list_id is None or list_kind != ordered:
                    require(depth == 0, '新列表必须从无缩进项开始；暂不支持混合嵌套列表')
                    list_id = list_numbering(self.document, ordered, int(marker[:-1]) if ordered else 1)
                    expected_numbers = {}
                    list_depth = 0
                require(depth <= list_depth + 1, '列表不能跳过缩进层级')
                if ordered:
                    current = int(marker[:-1])
                    require(current == expected_numbers.get(depth, current if depth == 0 else 1),
                            '有序列表编号必须连续；子列表从 1 开始')
                    expected_numbers[depth] = current + 1
                    expected_numbers = {d: n for d, n in expected_numbers.items() if d <= depth}
                paragraph = self.paragraph(text)
                numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                numbering.get_or_add_ilvl().val = depth
                numbering.get_or_add_numId().val = list_id
                list_kind, list_depth = ordered, depth
                continue
            else:
                require(not raw.startswith('    ') and not re.match(r'^(?:#{7,}|>|```|~~~|[-=]{3,}$|\[[^]]+\]:|\$\$)', line),
                        f'不支持的 Markdown 块语法：{line[:100]}')
                require(not re.fullmatch(r'(?:\*\s*){3,}|(?:_\s*){3,}', line), 'Markdown 分隔线暂不支持')
                paragraph_lines = [line]
                while index < len(lines) and lines[index].strip():
                    next_raw, next_line = lines[index], lines[index].strip()
                    if HEADING.match(next_line) or LIST.match(next_raw) or next_line.startswith(('!', '|', '表：', '>', '```', '~~~', '$$', '#')) or re.match(r'^\[[^]]+\]:', next_line):
                        break
                    require(not next_raw.startswith('    ') and not re.fullmatch(r'[-=]{3,}', next_line),
                            '缩进代码和 Setext 标题暂不支持')
                    require('\t' not in next_raw and not next_raw.endswith(('  ', '\\')),
                            '制表符或 Markdown 硬换行暂不支持，请使用独立段落')
                    paragraph_lines.append(next_line)
                    index += 1
                self.paragraph(' '.join(paragraph_lines))
            list_id = None


def export(workspace, output, template=None, profile_path=None, draft=False):
    """验证输入后原子写入 DOCX；模板只替换唯一独立正文锚点。"""
    workspace, output = Path(workspace).resolve(), Path(output).resolve()
    require(output.suffix.lower() == '.docx', '输出文件必须使用 .docx 扩展名')
    if (workspace / 'project.json').exists():
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        from project import build
        build(workspace)
    book, profile = load_book(workspace, draft), load_profile(profile_path)
    document = Document(str(template)) if template else Document()
    directories = toc_instructions(document)
    anchor = None
    if template:
        require(Path(template).resolve() != output, '输出不能覆盖原模板')
        validate_template(document, directories)
        anchors = [p for p in document.paragraphs if p.text.strip() == '{{正文}}']
        require(len(anchors) == 1, '模板必须含唯一独立 {{正文}} 段落；表格、文本框或未知模板不能自动填充')
        require(sum('{{正文}}' in ''.join(t.text or '' for t in node.findall('.//' + qn('w:t')))
                    for node in document.element.body.findall('.//' + qn('w:p'))) == 1,
                '模板含多个或非独立正文锚点，请清理后重试')
        anchor = anchors[0]
        require(anchor._p.find('.//' + qn('w:sectPr')) is None, '正文锚点不能同时承载分节符')
    configure(document, profile, bool(template))
    writer = Writer(document, workspace, profile, anchor)
    if draft:
        warning = writer.paragraph('草稿 未完成正式交付审核', 'Bid TOC Title')
        warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
        warning.paragraph_format.keep_with_next = True
    if not template:
        title = writer.paragraph(book['title'], 'Bid Cover')
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_before = Pt(100)
        title.paragraph_format.space_after = Pt(70)
        for value in (book['bidder'], book['date']):
            if value.strip():
                writer.paragraph(value).alignment = WD_ALIGN_PARAGRAPH.CENTER
        writer.paragraph().add_run().add_break(WD_BREAK.PAGE)
    if not directories:
        writer.paragraph('目录', 'Bid TOC Title')
        field(writer.paragraph(), f'TOC \\o "1-{profile["toc_levels"]}" \\h \\z \\u', '请在 Word 中更新整个目录')
        writer.paragraph().add_run().add_break(WD_BREAK.PAGE)
    for index, section in enumerate(book['sections']):
        writer.heading(section['number'], section['title'], section['level'])
        has_children = index + 1 < len(book['sections']) and book['sections'][index + 1]['level'] > section['level']
        writer.markdown(section, has_children)
    if anchor is not None:
        anchor._p.getparent().remove(anchor._p)
    document.core_properties.title = book['title']
    document.core_properties.author = book['bidder']
    output.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(suffix='.docx', prefix='.bid-', dir=output.parent)
    os.close(handle)
    try:
        document.save(temporary)
        os.replace(temporary, output)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return {'output': str(output), 'draft': draft, 'tables': writer.tables, 'images': writer.images,
            'next': '在 Word 或 LibreOffice 更新全部字段，并逐页检查排版；目录页码由排版软件生成。'}


def main():
    """提供可直接供 WorkBuddy 调用的命令行入口。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workspace')
    parser.add_argument('--output', required=True)
    parser.add_argument('--template')
    parser.add_argument('--profile')
    parser.add_argument('--draft', action='store_true')
    args = parser.parse_args()
    try:
        result = export(args.workspace, args.output, args.template, args.profile, args.draft)
    except (ValueError, OSError, KeyError) as error:
        parser.exit(2, f'导出失败：{error}\n')
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    main()

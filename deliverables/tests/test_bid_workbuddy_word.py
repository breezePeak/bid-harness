"""验证 Word 导出的实际 OOXML、模板保留、Markdown 拒绝与正式交付条件。"""

import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
import subprocess
import sys
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from PIL import Image


SCRIPT = Path(__file__).resolve().parents[1] / 'bid-workbuddy' / 'scripts' / 'export_word.py'
SPEC = importlib.util.spec_from_file_location('export_word', SCRIPT)
export_word = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(export_word)


class WordExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        (self.root / 'build').mkdir()
        self.output = self.root / 'output.docx'
        self.book = {'title': '设备运维技术方案', 'bidder': '测试投标单位', 'date': '2026年10月',
                     'status': 'ready', 'issues': [], 'sections': [
                         {'title': '总体方案', 'number': '1', 'level': 1, 'body': '正文内容。'}]}

    def save(self):
        (self.root / 'build' / 'book.json').write_text(json.dumps(self.book, ensure_ascii=False), encoding='utf-8')

    def export(self, **kwargs):
        self.save()
        return export_word.export(self.root, self.output, **kwargs)

    def xml(self, part='word/document.xml'):
        with ZipFile(self.output) as archive:
            return archive.read(part).decode('utf-8')

    def test_chinese_a4_cover_real_toc_and_page_fields(self):
        self.export()
        document = Document(self.output)
        self.assertAlmostEqual(document.sections[0].page_width.cm, 21, places=2)
        self.assertAlmostEqual(document.sections[0].page_height.cm, 29.7, places=2)
        self.assertEqual(document.paragraphs[0].text, self.book['title'])
        self.assertIn('TOC \\o "1-3" \\h \\z \\u', self.xml())
        self.assertIn(' PAGE ', self.xml('word/footer1.xml'))
        self.assertNotIn('NUMPAGES', self.xml())
        self.assertIn('w:updateFields w:val="true"', self.xml('word/settings.xml'))
        self.assertIn('w:eastAsia="宋体"', self.xml('word/styles.xml'))
        self.assertNotRegex(self.xml(), r'总体方案[.·…]+[0-9]')

    def test_hierarchy_and_numbered_body_headings(self):
        self.book['sections'] = [
            {'title': '总体方案', 'number': '1', 'level': 1, 'body': '概述。'},
            {'title': '实施方法', 'number': '1.1', 'level': 2, 'body': '# 组织方法\n\n## 责任划分\n\n具体正文。\n\n# 验收安排'},
            {'title': '保障措施', 'number': '2', 'level': 1, 'body': ''},
        ]
        self.export()
        headings = [(p.text, p.style.name) for p in Document(self.output).paragraphs if 'Bid Heading' in p.style.name]
        self.assertIn(('1.1.1 组织方法', 'Bid Heading 3'), headings)
        self.assertIn(('1.1.1.1 责任划分', 'Bid Heading 4'), headings)
        self.assertIn(('1.1.2 验收安排', 'Bid Heading 3'), headings)

    def test_tables_repeat_headers_width_and_images(self):
        (self.root / 'assets').mkdir()
        Image.new('RGB', (2400, 800), 'white').save(self.root / 'assets' / '流程.png')
        self.book['sections'][0]['body'] = (
            '**重点措施**与*辅助说明*。\n\n'
            '- 第一项\n- 第二项\n\n1. 第一步\n2. 第二步\n\n'
            '表：责任清单\n| 角色 | 职责 |\n| :---: | ---: |\n| 负责人 | 组织实施与质量复核 |\n\n'
            '![实施流程](assets/流程.png)\n\n[资料来源](https://example.com/source)')
        result = self.export()
        document = Document(self.output)
        self.assertEqual((result['tables'], result['images']), (1, 1))
        self.assertEqual(document.tables[0].cell(1, 1).text, '组织实施与质量复核')
        self.assertIn('w:tblHeader', self.xml())
        self.assertIn('w:numPr', self.xml())
        self.assertIn('w:hyperlink', self.xml())
        self.assertEqual(document.tables[0].cell(1, 1).paragraphs[0].alignment, 2)
        self.assertIn('表 1 责任清单', '\n'.join(p.text for p in document.paragraphs))
        self.assertLessEqual(document.inline_shapes[0].width, document.sections[0].page_width
                             - document.sections[0].left_margin - document.sections[0].right_margin)
        self.assertEqual(document.inline_shapes[0]._inline.docPr.get('descr'), '实施流程')
        self.assertTrue(any(r.bold and r.text == '重点措施' for p in document.paragraphs for r in p.runs))

    def test_template_anchor_keeps_surroundings_and_page_setup(self):
        template = self.root / 'template.docx'
        document = Document()
        document.sections[0].page_width = Cm(20)
        document.add_paragraph('模板封面内容')
        document.add_paragraph('{{正文}}')
        document.add_table(rows=1, cols=1).cell(0, 0).text = '模板原有签章表'
        document.add_paragraph('模板结束文字')
        document.sections[0].header.paragraphs[0].text = '保留页眉'
        document.save(template)
        self.export(template=template)
        document = Document(self.output)
        text = '\n'.join(p.text for p in document.paragraphs)
        self.assertNotIn('{{正文}}', text)
        self.assertTrue(text.startswith('模板封面内容'))
        self.assertTrue(text.endswith('模板结束文字'))
        self.assertLess(text.index('正文内容'), text.index('模板结束文字'))
        self.assertEqual(document.tables[0].cell(0, 0).text, '模板原有签章表')
        self.assertEqual(document.sections[0].header.paragraphs[0].text, '保留页眉')
        self.assertAlmostEqual(document.sections[0].page_width.cm, 20, places=2)

    def test_ambiguous_template_rejected_without_overwriting_output(self):
        self.output.write_bytes(b'previous result')
        for anchors in ([], ['{{正文}}', '{{正文}}'], ['前缀{{正文}}']):
            with self.subTest(anchors=anchors):
                template = self.root / 'bad-template.docx'
                document = Document()
                for text in anchors:
                    document.add_paragraph(text)
                document.save(template)
                with self.assertRaisesRegex(ValueError, '唯一独立'):
                    self.export(template=template)
                self.assertEqual(self.output.read_bytes(), b'previous result')

    def test_draft_requires_flag_and_marks_document(self):
        self.book['status'] = 'draft'
        self.book['issues'] = ['存在待核实的企业事实']
        with self.assertRaisesRegex(ValueError, '尚未 ready'):
            self.export()
        self.export(draft=True)
        self.assertIn('草稿 未完成正式交付审核', self.xml())

    def test_unknown_cover_metadata_is_omitted(self):
        self.book['bidder'], self.book['date'] = '', ''
        self.export()
        self.assertNotIn('待补', self.xml())

    def test_real_project_rebuild_rejects_stale_ready_cache(self):
        project_root = self.root / 'project'
        subprocess.run([sys.executable, '-B', str(SCRIPT.with_name('project.py')), 'init', str(project_root),
                        '--title', '实际未完成项目'], check=True, capture_output=True)
        (project_root / 'build').mkdir()
        (project_root / 'build' / 'book.json').write_text(json.dumps(self.book, ensure_ascii=False), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '尚未 ready'):
            export_word.export(project_root, self.output)
        refreshed = json.loads((project_root / 'build' / 'book.json').read_text(encoding='utf-8'))
        self.assertEqual(refreshed['status'], 'draft')
        self.assertEqual(refreshed['title'], '实际未完成项目')
        self.assertFalse(self.output.exists())

    def test_existing_template_toc_is_preserved_once(self):
        template = self.root / 'template-toc.docx'
        document = Document()
        export_word.field(document.add_paragraph(), 'TOC \\o "1-6" \\h \\z \\u', '原模板目录')
        document.add_paragraph('{{正文}}')
        document.save(template)
        self.export(template=template)
        self.assertEqual(self.xml().count(' TOC '), 1)
        self.assertIn('原模板目录', self.xml())

    def test_multi_column_template_rejected_before_output(self):
        template = self.root / 'two-columns.docx'
        for explicit_columns in (False, True):
            with self.subTest(explicit_columns=explicit_columns):
                document = Document()
                document.add_paragraph('{{正文}}')
                columns = document.sections[0]._sectPr.find(qn('w:cols'))
                if explicit_columns:
                    columns.append(export_word.element('col', w=3000))
                    columns.append(export_word.element('col', w=3000))
                else:
                    columns.set(qn('w:num'), '2')
                document.save(template)
                original = template.read_bytes()
                with self.assertRaisesRegex(ValueError, '多栏模板暂不支持'):
                    self.export(template=template)
                self.assertFalse(self.output.exists())
                self.assertEqual(template.read_bytes(), original)

    def test_custom_style_only_template_toc_rejected(self):
        template = self.root / 'custom-toc.docx'
        for simple_field in (False, True):
            with self.subTest(simple_field=simple_field):
                document = Document()
                paragraph = document.add_paragraph()
                code = ' TOC \\t "客户章节,1" \\h '
                if simple_field:
                    paragraph._p.append(export_word.element('fldSimple', instr=code))
                else:
                    export_word.field(paragraph, code, '客户目录')
                    instruction = paragraph._p.find('.//' + qn('w:instrText'))
                    instruction.text = ' TO'
                    split_run = export_word.element('r')
                    split_text = export_word.element('instrText')
                    split_text.text = 'C \\t "客户章节,1" \\h '
                    split_run.append(split_text)
                    instruction.getparent().addnext(split_run)
                document.add_paragraph('{{正文}}')
                document.save(template)
                original = template.read_bytes()
                with self.assertRaisesRegex(ValueError, '模板原有目录未启用大纲级别'):
                    self.export(template=template)
                self.assertFalse(self.output.exists())
                self.assertEqual(template.read_bytes(), original)

    def test_simple_template_toc_with_outline_option_preserved(self):
        template = self.root / 'simple-toc.docx'
        document = Document()
        instruction = ' TOC \\o "1-3" \\u '
        document.add_paragraph()._p.append(export_word.element('fldSimple', instr=instruction))
        document.add_paragraph('{{正文}}')
        document.save(template)
        self.export(template=template)
        self.assertEqual(export_word.toc_instructions(Document(self.output)), [instruction])

    def test_images_reject_network_absolute_and_escaping_paths(self):
        for path in ('https://example.com/a.png', '../outside.png', '%2e%2e/outside.png', 'C:/a.png', '/tmp/a.png'):
            with self.subTest(path=path):
                self.book['sections'][0]['body'] = f'![图片]({path})'
                with self.assertRaisesRegex(ValueError, '图片'):
                    self.export()

    def test_unsupported_markdown_is_not_silently_lost(self):
        for body in ('```python\nprint(1)\n```', '> 引用块', '文本 `代码`', '![无图]()',
                     '- [ ] 任务项', '<div>HTML</div>', '正文\n===', '~~删除线~~',
                     '| A | B |\n| --- | --- |\n| 少一列 |', '**未闭合', '硬换行  \n下一行',
                     '段落\n[x]: https://example.com', '文字 $x+y$', '正文\n下一行\t缩进'):
            with self.subTest(body=body):
                self.book['sections'][0]['body'] = body
                with self.assertRaises(ValueError):
                    self.export()

    def test_deep_body_heading_rejected(self):
        self.book['sections'] = [
            {'title': f'第{level}层', 'number': '.'.join(['1'] * level), 'level': level,
             'body': '# 超过六级' if level == 6 else ''} for level in range(1, 7)]
        with self.assertRaisesRegex(ValueError, '超过 6'):
            self.export()

    def test_invalid_heading_number_and_level_rejected(self):
        for number, level in [('2', 1), ('1.1', 2), ('1', 7)]:
            with self.subTest(number=number, level=level):
                self.book['sections'][0].update(number=number, level=level)
                with self.assertRaisesRegex(ValueError, '层级|编号'):
                    self.export()

    def test_parent_section_cannot_duplicate_child_headings(self):
        self.book['sections'][0]['body'] = '# 私自增加的标题'
        self.book['sections'].append({'title': '正式子节', 'number': '1.1', 'level': 2, 'body': '内容'})
        with self.assertRaisesRegex(ValueError, '正式下级章节'):
            self.export()

    def test_profile_unknown_and_out_of_range_values_rejected(self):
        path = self.root / 'profile.json'
        for profile in ({'fake_page_count': 100}, {'body_size_pt': 2}, {'toc_levels': 7}, {'margin_left_cm': True}):
            with self.subTest(profile=profile):
                path.write_text(json.dumps(profile), encoding='utf-8')
                with self.assertRaises(ValueError):
                    self.export(profile_path=path)


if __name__ == '__main__':
    unittest.main()

"""标书技能从文件接入到 DOCX 及 HTML 预览的阶段门禁回归。"""

import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile


SCRIPT = Path(os.environ.get('BID_WRITING_WORKFLOW_SCRIPT') or
              Path(__file__).resolve().parents[1] / 'bid-writing-workflow' / 'scripts' / 'workflow.py')


def call(*args, okay=True):
    result = subprocess.run([sys.executable, str(SCRIPT), *map(str, args)], capture_output=True,
                            text=True, encoding='utf-8', env={**os.environ, 'PYTHONIOENCODING': 'utf-8'})
    if okay and result.returncode:
        raise AssertionError(result.stderr)
    return result


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
    return path


class WorkflowTest(unittest.TestCase):
    def test_full_run_and_upstream_revision(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as temporary:
            base = Path(temporary)
            tender = base / '招标文件.txt'
            tender.write_text('采购实施服务。技术要求：实施计划。评分：实施步骤与进度。', encoding='utf-8')
            work = base / '标书'
            result = json.loads(call('init', work, '--tender', tender).stdout)
            intake = Path(result['candidate'])
            self.assertNotEqual(call('confirm', work, 'file_intake', '--user-confirmation', '确认', okay=False).returncode, 0)
            intake_data = json.loads(intake.read_text(encoding='utf-8'))
            intake_data['files'][0]['readable'] = True
            intake_data['files'][0].pop('issue')
            intake_data['inspection'] = '已核对正文和评分，附件无。'
            write(intake, intake_data)
            source = intake_data['files'][0]['path']
            submitted = json.loads(call('submit', work, 'file_intake', intake).stdout)
            first_preview = Path(submitted['preview'])
            self.assertTrue(first_preview.is_file())
            self.assertIn('文件接入', first_preview.read_text(encoding='utf-8'))
            self.assertNotEqual(call('submit', work, 'tender_analysis', intake, okay=False).returncode, 0)
            call('confirm', work, 'file_intake', '--user-confirmation', '确认文件接入结果')
            self.assertIn('确认文件接入结果', first_preview.read_text(encoding='utf-8'))

            analysis = {
                'project': {'title': '示例技术标', 'summary': '采购实施服务<script>alert(1)</script>'},
                'requirements': [{'text': '实施计划', 'mandatory': True,
                                  'source': {'file': source, 'location': '第1条'}}],
                'scoring': [{'title': '实施方案', 'original_text': '步骤与进度完整得分',
                             'source': {'file': source, 'location': '评分表第1行'}}],
                'compliance': [],
            }
            analysis_path = write(base / 'analysis.json', analysis)
            analysis_preview = Path(json.loads(call('submit', work, 'tender_analysis', analysis_path).stdout)['preview'])
            self.assertIn('&lt;script&gt;alert(1)&lt;/script&gt;', analysis_preview.read_text(encoding='utf-8'))
            self.assertNotIn('<script>alert(1)</script>', analysis_preview.read_text(encoding='utf-8'))
            call('confirm', work, 'tender_analysis', '--user-confirmation', '确认招标分析')

            section = {'title': '实施方案', 'parent_position': None, 'writable': True,
                       'must_answer': ['如何实施并安排进度'], 'requirement_positions': [1],
                       'scoring_positions': [1], 'response_point_positions': [1, 2],
                       'compliance_positions': []}
            outline = {'response_points': [['实施步骤', '进度安排']], 'sections': [section],
                       'global_compliance_positions': []}
            bad = {**outline, 'sections': [{**section, 'response_point_positions': [1]}]}
            self.assertIn('响应点 2', call('submit', work, 'initial_outline', write(base / 'bad.json', bad), okay=False).stderr)
            outline_path = write(base / 'outline.json', outline)
            outline_preview = Path(json.loads(call('submit', work, 'initial_outline', outline_path).stdout)['preview'])
            self.assertIn('评分响应点', outline_preview.read_text(encoding='utf-8'))
            call('confirm', work, 'initial_outline', '--user-confirmation', '确认初版目录')

            final = {**outline, 'research': [{'section_position': 1, 'evidence': [
                {'file': source, 'location': '第1条', 'use': '实施范围依据'}],
                'missing_topics': ['企业人员待核实'], 'answer_plan': '写步骤和进度。'}]}
            evidence_preview = Path(json.loads(call('submit', work, 'evidence_mapping', write(base / 'final.json', final)).stdout)['preview'])
            self.assertIn('企业人员待核实', evidence_preview.read_text(encoding='utf-8'))
            call('confirm', work, 'evidence_mapping', '--user-confirmation', '确认终版目录')

            confirmed_outline = work / 'artifacts' / '04-final-outline.json'
            original = confirmed_outline.read_bytes()
            confirmed_outline.write_bytes(original + b' ')
            self.assertIn('已确认产物发生改变', call('submit', work, 'chapter_writing',
                                                     write(base / 'empty.json', {}), okay=False).stderr)
            confirmed_outline.write_bytes(original)

            manuscript = {'chapters': [{'section_position': 1, 'body': '我方拟分阶段实施。\n\n| 阶段 | 成果 |\n| --- | --- |\n| 启动 | 计划 |',
                                        'answered_requirement_positions': [1],
                                        'answered_response_point_positions': [1, 2], 'review': 'passed'}],
                          'whole_book_review': {'result': 'passed', 'issues': []}}
            manuscript_preview = Path(json.loads(call('submit', work, 'chapter_writing', write(base / 'manuscript.json', manuscript)).stdout)['preview'])
            self.assertIn('我方拟分阶段实施', manuscript_preview.read_text(encoding='utf-8'))
            call('confirm', work, 'chapter_writing', '--user-confirmation', '确认全书正文')
            exported = json.loads(call('export', work).stdout)
            export_preview = Path(exported['preview'])
            self.assertIn('打开或下载 DOCX', export_preview.read_text(encoding='utf-8'))
            with ZipFile(exported['file']) as archive:
                self.assertEqual(archive.testzip(), None)
                document = archive.read('word/document.xml').decode('utf-8')
                self.assertIn('我方拟分阶段实施', document)
                self.assertIn('w:tbl', document)
            try:
                from docx import Document
            except ImportError:
                pass
            else:
                opened = Document(exported['file'])
                self.assertIn('实施方案', [paragraph.text for paragraph in opened.paragraphs])
                self.assertEqual(len(opened.tables), 1)
            self.assertNotEqual(call('submit-docx', work, tender, okay=False).returncode, 0)
            self.assertEqual(json.loads(call('submit-docx', work, exported['file']).stdout)['preview'], str(export_preview))
            call('confirm', work, 'docx_export', '--user-confirmation', '确认 Word 版面')
            self.assertIsNone(json.loads(call('status', work).stdout)['current_stage'])

            analysis['project']['summary'] = '修改后的采购实施范围'
            call('submit', work, 'tender_analysis', write(analysis_path, analysis))
            status = json.loads(call('status', work).stdout)
            self.assertEqual(status['current_stage'], 'tender_analysis')
            self.assertEqual([item['status'] for item in status['stages'][2:]], ['pending'] * 4)
            self.assertTrue(analysis_preview.is_file())
            self.assertFalse(outline_preview.exists())
            self.assertFalse(evidence_preview.exists())
            self.assertFalse(manuscript_preview.exists())
            self.assertFalse(export_preview.exists())
            self.assertNotEqual(call('export', work, okay=False).returncode, 0)

    def test_adding_source_resets_intake_confirmation(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as temporary:
            base = Path(temporary)
            tender = base / 'tender.txt'
            tender.write_text('技术要求', encoding='utf-8')
            work = base / 'work'
            candidate = Path(json.loads(call('init', work, '--tender', tender).stdout)['candidate'])
            data = json.loads(candidate.read_text(encoding='utf-8'))
            data['files'][0]['readable'] = True
            data['inspection'] = '已核对正文。'
            write(candidate, data)
            preview = Path(json.loads(call('submit', work, 'file_intake', candidate).stdout)['preview'])
            call('confirm', work, 'file_intake', '--user-confirmation', '确认资料')
            reference = base / 'reference.txt'
            reference.write_text('参考资料', encoding='utf-8')
            call('add-source', work, '--role', 'reference', '--file', reference)
            status = json.loads(call('status', work).stdout)
            self.assertEqual(status['current_stage'], 'file_intake')
            self.assertTrue(all(item['status'] == 'pending' for item in status['stages']))
            self.assertEqual(len(json.loads(candidate.read_text(encoding='utf-8'))['files']), 2)
            self.assertFalse(preview.exists())


if __name__ == '__main__':
    unittest.main()
